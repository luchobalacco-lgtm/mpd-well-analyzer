from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import FileResponse, HTMLResponse
import pandas as pd
import openpyxl
import re, io, tempfile

app = FastAPI()

def read_time_summary(xl_bytes):
    df = pd.read_excel(io.BytesIO(xl_bytes), sheet_name="Time Summary", header=None)
    df.columns = ["_","Pozo","Equipo","EventCode","Desde","Hasta","Horas","Fase","Tarea","Actividad","Codigo","NPT","Detalle_NPT","Evidencia","CausaRaiz","Prof","Prof_NPT","Operacion","Descripcion","Equipment","FailedEquip","ServiceCompany"]
    df = df.iloc[1:].copy()
    df["Desde"] = pd.to_datetime(df["Desde"], errors="coerce")
    df["Hasta"] = pd.to_datetime(df["Hasta"], errors="coerce")
    df["Horas"] = pd.to_numeric(df["Horas"], errors="coerce")
    return df

def read_sbp(xl_bytes):
    xl = pd.ExcelFile(io.BytesIO(xl_bytes))
    sbp_name = next((s for s in xl.sheet_names if s.upper() in ("SBP","SPB")), None)
    if not sbp_name:
        return None
    df = pd.read_excel(io.BytesIO(xl_bytes), sheet_name=sbp_name, header=None)
    df.columns = ["_","Time","MD","TVD","FlowIn","CoriolisRate","ReturnFlow","SBP_psi","RCD_psi","CoriolicsMW","ROP","RPM","ECD","ChokeA","ChokeB","Estado","Well"]
    df = df.iloc[1:].copy()
    df["Time"] = pd.to_datetime(df["Time"], errors="coerce")
    df["RPM"] = pd.to_numeric(df["RPM"].astype(str).str.replace(",","."), errors="coerce").fillna(0)
    return df

def read_bearing_serial(xl_bytes):
    xl = pd.ExcelFile(io.BytesIO(xl_bytes))
    bs_name = next((s for s in xl.sheet_names if any(w in s.lower() for w in ["bearing","serial","hoja"])), None)
    if not bs_name:
        return None
    df = pd.read_excel(io.BytesIO(xl_bytes), sheet_name=bs_name, header=None)
    df = df.iloc[1:].copy()
    df.columns = ["_","Well","InstallDate","PullDate","RunNo","Serial"]
    df["InstallDate"] = pd.to_datetime(df["InstallDate"], errors="coerce")
    df["PullDate"] = pd.to_datetime(df["PullDate"], errors="coerce")
    df["RunNo"] = pd.to_numeric(df["RunNo"], errors="coerce")
    return df.dropna(subset=["InstallDate","Serial"])

def get_well_name(ts):
    vals = ts["Pozo"].dropna().unique()
    return str(vals[0]) if len(vals) > 0 else "Unknown"

def find_service_start(ts):
    bopsur = ts[ts["Tarea"] == "BOPSUR"].sort_values("Desde")
    if bopsur.empty:
        return ts.iloc[0]["Desde"], 0.0
    mpd_kw = r"MPD|campana de viaje|4000 PSI|4\.000 PSI|RCD"
    mpd_rows = bopsur[bopsur["Operacion"].str.contains(mpd_kw, case=False, na=False)]
    if not mpd_rows.empty:
        target_phase = mpd_rows.iloc[0]["Fase"]
        t_ref = mpd_rows.iloc[0]["Desde"]
        phase_rows = bopsur[(bopsur["Fase"] == target_phase) & (bopsur["Desde"] >= t_ref - pd.Timedelta(hours=24)) & (bopsur["Desde"] <= t_ref + pd.Timedelta(hours=24))]
        return phase_rows["Hasta"].max(), round(float(phase_rows["Horas"].sum()), 2)
    return bopsur.iloc[0]["Hasta"], round(float(bopsur["Horas"].sum()), 2)

def find_npt_mpd(ts, start_date):
    period = ts[ts["Desde"] >= start_date].copy()
    combos = {"WREP-RMPD": {"hs": 0.0, "detail": ""}, "WSER-SMPD": {"hs": 0.0, "detail": ""}, "ROT-SFAL-RTME-MPD": {"hs": 0.0, "detail": ""}, "ROT-LOSS-CECD": {"hs": 0.0, "detail": ""}, "ROT-WCON-CECD": {"hs": 0.0, "detail": ""}}
    npt_rows = period[period["NPT"].notna()].copy()
    if npt_rows.empty:
        return combos
    npt_rows["npt_base"] = npt_rows["NPT"].astype(str).str.strip() + "-" + npt_rows["Detalle_NPT"].astype(str).str.strip()
    for combo in ["WREP-RMPD", "WSER-SMPD", "ROT-SFAL-RTME-MPD"]:
        matched = npt_rows[npt_rows["npt_base"].str.upper() == combo.upper()]
        if not matched.empty:
            combos[combo]["hs"] = round(float(matched["Horas"].sum()), 2)
            combos[combo]["detail"] = " | ".join(matched["Operacion"].dropna().astype(str).str[:200].tolist())
    for col in ["Evidencia", "CausaRaiz"]:
        if col not in npt_rows.columns:
            continue
        cecd_rows = npt_rows[npt_rows[col].astype(str).str.contains("CECDMPD", case=False, na=False)]
        if cecd_rows.empty:
            continue
        for base, target in [("ROT-WCON", "ROT-WCON-CECD"), ("ROT-LOSS", "ROT-LOSS-CECD")]:
            matched = cecd_rows[cecd_rows["npt_base"].str.upper() == base.upper()]
            if not matched.empty:
                combos[target]["hs"] = round(float(matched["Horas"].sum()), 2)
                combos[target]["detail"] = " | ".join(matched["Operacion"].dropna().astype(str).str[:200].tolist())
    return combos

DEPTH_RE = re.compile(r'(?:DESDE|DE)\s+([\d,\.]+)\s*[mM]?\s+(?:HASTA|A)\s+([\d,\.]+)\s*[mM]', re.IGNORECASE)
SAFETY_KW = re.compile(r"REUNI|HSE|CHARLA|SEGURIDAD|OPERATIVA|PREVIO|FLOW CHECK|FC PREVIO", re.IGNORECASE)
INSTALL_RE = re.compile(r"(COLOC[AO][NR]?\s+BEARING|COLOC[AO][NR]?\s+BERING|MONT[AO]\s+BEARING|MONT[AO]\s+BERING|INSTALA\w*\s+BEARING|INSTALA\w*\s+BERING|INSTALA\w*\s+EN\s+RCD|RETIRAD?\w*\s+CAMPANA.*BEARING|RETIRAD?\w*\s+CAMPANA.*BERING|RETIRO\s+DE\s+CAMPANA.*BEARING|RETIRO\s+DE\s+CAMPANA.*BERING)", re.IGNORECASE)
REMOVE_RE = re.compile(r"(RETIRA\w*\s+BEARING|RETIRA\w*\s+BERING|DESMONTA\w*\s+BEARING|SACA\w*\s+BEARING|CAMBIA\w*\s+BEARING|REEMPLAZA\w*\s+BEARING)", re.IGNORECASE)

def calc_stripped_meters(seg):
    strip_rows = seg[seg["Actividad"].isin(["POH","RIH","TBHNDL"]) & ~seg["Operacion"].str.contains(r"CHARLA|SEGURIDAD|PREVIO|FLOW CHECK|ESPERA|REUNI", case=False, na=False)].sort_values("Desde")
    if strip_rows.empty:
        return 0
    pairs = []
    for _, row in strip_rows.iterrows():
        for a, b in DEPTH_RE.findall(str(row["Operacion"])):
            d1, d2 = float(a.replace(",",".")), float(b.replace(",","."))
            if d1 > 100 and d2 > 100:
                pairs.append((d1, d2, row["Actividad"]))
    if not pairs:
        return 0
    total = 0
    max_down = None
    min_up = None
    for d1, d2, act in pairs:
        lo, hi = min(d1,d2), max(d1,d2)
        if act in ("RIH","TBHNDL"):
            if max_down is None:
                total += hi - lo
                max_down = hi
            elif hi > max_down:
                total += hi - max(max_down, lo)
                max_down = hi
        else:
            if min_up is None:
                total += hi - lo
                min_up = lo
            elif lo < min_up:
                total += min(min_up, hi) - lo
                min_up = lo
    return round(total)

def classify_bearing_row(op_text):
    if SAFETY_KW.search(op_text):
        return None
    has_install = bool(INSTALL_RE.search(op_text))
    has_remove = bool(REMOVE_RE.search(op_text))
    if has_install and not has_remove:
        return "install"
    if has_remove and not has_install:
        return "remove"
    if has_install and has_remove:
        return "install" if INSTALL_RE.search(op_text).start() < REMOVE_RE.search(op_text).start() else "remove"
    return None

def find_bearings(ts, sbp, serial_df, start_date):
    period = ts[ts["Desde"] >= start_date].copy().sort_values("Desde")
    mask = period["Operacion"].str.contains(r"BEARING|BERING", case=False, na=False)
    b_rows = period[mask].reset_index(drop=True)
    installs, removes = [], []
    for _, row in b_rows.iterrows():
        kind = classify_bearing_row(str(row["Operacion"]))
        if kind == "install":
            installs.append(row)
        elif kind == "remove":
            removes.append(row)
    pairs = []
    used = set()
    for inst in installs:
        for i, rem in enumerate(removes):
            if i in used:
                continue
            if rem["Hasta"] > inst["Desde"]:
                used.add(i)
                pairs.append({"in_row": inst, "out_row": rem})
                break
    last_idx = len(pairs) - 1
    results = []
    for idx, pair in enumerate(pairs):
        t_in, t_out = pair["in_row"]["Desde"], pair["out_row"]["Hasta"]
        seg = period[(period["Desde"] >= t_in) & (period["Hasta"] <= t_out)]
        drill = seg[seg["Actividad"] == "DRL"]
        total_drill = 0.0
        for _, row in drill.iterrows():
            m = re.search(r"AVANCE:\s*(\d+(?:[.,]\d+)?)\s*m", str(row["Operacion"]), re.IGNORECASE)
            if m:
                total_drill += float(m.group(1).replace(",","."))
        strip_m = calc_stripped_meters(seg)
        svc = round((t_out - t_in).total_seconds() / 3600, 2)
        rot_hs = 0
        if sbp is not None:
            s = sbp[(sbp["Time"] >= t_in) & (sbp["Time"] <= t_out)]
            rot_hs = int((s["RPM"] > 0).sum())
        op_out = str(pair["out_row"]["Operacion"]).upper()
        motivo = "Campana de Viaje"
        if any(w in op_out for w in ["CEMENT","CEMEN"]):
            motivo = "Cementación"
        elif any(w in op_out for w in ["DAÑ","FALLA","LEAK","FUGA","ROTURA"]):
            motivo = "Daño de Bearing"
        elif "ELASTOM" in op_out:
            motivo = "Daño de Elastomero"
        elif "HORAS" in op_out:
            motivo = "Horas Acumuladas"
        elif idx == last_idx:
            after = period[period["Desde"] >= t_out].head(10)
            for _, row in after.iterrows():
                if any(w in str(row["Operacion"]).upper() for w in ["CEMENT","CEMEN","CIA CEMENTACION","CIA CEMENTACIÓN"]):
                    motivo = "Cementación"
                    break
        tcambio = 0.0
        out_op = str(pair["out_row"]["Operacion"])
        if not SAFETY_KW.search(out_op):
            tcambio = round(float(pair["out_row"]["Horas"]) if pd.notna(pair["out_row"]["Horas"]) else 0.0, 2)
        presion = "SIN"
        if any(w in op_out for w in ["PSI","PRESION","BOP CERR","ANULAR CERR","STRIPPING","INCREMENTA"]):
            presion = "CON"
        serial = "No cargado en OW"
        if serial_df is not None:
            matched = serial_df[(serial_df["InstallDate"] >= t_in - pd.Timedelta(hours=3)) & (serial_df["InstallDate"] <= t_in + pd.Timedelta(hours=3))]
            if not matched.empty:
                serial = str(matched.iloc[0]["Serial"])
        results.append({"run": idx+1, "serial": serial, "t_in": t_in, "t_out": t_out, "drill_m": round(total_drill), "strip_m": strip_m, "svc_hs": svc, "motivo": motivo, "tcambio": tcambio, "rot_hs": rot_hs, "presion": presion})
    return results

def fill_template(template_bytes, well_name, inicio, fin, total_hs, npt_combos, bearings):
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    if "Horas Operativas del Servicio" in wb.sheetnames:
        ws = wb["Horas Operativas del Servicio"]
        ws["B3"] = well_name; ws["C3"] = "SLB"; ws["D3"] = inicio; ws["D3"].number_format = "DD/MM/YYYY HH:MM"; ws["E3"] = fin; ws["E3"].number_format = "DD/MM/YYYY HH:MM"; ws["F3"] = total_hs
    if "NPT" in wb.sheetnames:
        ws = wb["NPT"]
        total = round(sum(v["hs"] for v in npt_combos.values()), 2)
        for i, (combo, data) in enumerate(npt_combos.items(), 3):
            ws[f"B{i}"] = well_name; ws[f"C{i}"] = "SLB"; ws[f"D{i}"] = combo; ws[f"E{i}"] = data["hs"]; ws[f"F{i}"] = total
            if data["detail"]: ws[f"G{i}"] = data["detail"]
    if "Bearing Resumen" in wb.sheetnames:
        ws = wb["Bearing Resumen"]
        for i, b in enumerate(bearings, 3):
            ws[f"B{i}"] = well_name; ws[f"C{i}"] = "SLB"; ws[f"D{i}"] = b["run"]; ws[f"E{i}"] = b["serial"]; ws[f"F{i}"] = b["drill_m"]
            ws[f"G{i}"] = b["strip_m"] if b["strip_m"] > 0 else "REVISAR"; ws[f"H{i}"] = b["svc_hs"]; ws[f"I{i}"] = b["motivo"]; ws[f"J{i}"] = b["tcambio"]; ws[f"K{i}"] = b["rot_hs"]; ws[f"L{i}"] = b["presion"]
    out = io.BytesIO(); wb.save(out); out.seek(0); return out.read()

@app.get("/", response_class=HTMLResponse)
async def index():
    with open("templates/index.html", "r", encoding="utf-8") as f:
        return f.read()

@app.post("/process")
async def process(data_file: UploadFile = File(...), template_file: UploadFile = File(...)):
    try:
        data_bytes = await data_file.read(); template_bytes = await template_file.read()
        ts = read_time_summary(data_bytes); sbp = read_sbp(data_bytes); serial_df = read_bearing_serial(data_bytes)
        well_name = get_well_name(ts); inicio, _ = find_service_start(ts)
        fin = ts.iloc[-1]["Hasta"]; total_hs = round((fin - inicio).total_seconds() / 3600, 2)
        npt_combos = find_npt_mpd(ts, inicio); bearings = find_bearings(ts, sbp, serial_df, inicio)
        result_bytes = fill_template(template_bytes, well_name, inicio, fin, total_hs, npt_combos, bearings)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx"); tmp.write(result_bytes); tmp.close()
        filename = f"{well_name.replace('(','').replace(')','').replace(' ','_')}_resultado.xlsx"
        return FileResponse(tmp.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=filename)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
